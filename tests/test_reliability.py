from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from unittest.mock import MagicMock

import httpx
import pytest

from deliverybrief.budget import BudgetLedger
from deliverybrief.demo_data import DEMO_PERIOD
from deliverybrief.exports import action_csv, approved_exports
from deliverybrief.generator import DemoReportGenerator
from deliverybrief.intake import normalize_evidence, parse_upload, scoped_evidence
from deliverybrief.integrations.github import GitHubEvidenceClient
from deliverybrief.integrations.google_docs import GoogleDocsEvidenceClient
from deliverybrief.models import ApprovalStatus, FindingSeverity, ProjectConfig
from deliverybrief.privacy import redact, safe_evidence
from deliverybrief.reliability_cases import CASES
from deliverybrief.report_documents import client_docx, client_pdf
from deliverybrief.scenarios import sample_item
from deliverybrief.storage import RunStore
from deliverybrief.validator import validate_report
from deliverybrief.workflow import generate_and_record


def generated(items):
    return DemoReportGenerator().generate(ProjectConfig(), DEMO_PERIOD, items).report


def setup_run(tmp_path):
    items = [
        sample_item("A", "Merged audit logging.", source="github"),
        sample_item("B", "Confirm acceptance. Owner: PM."),
    ]
    store = RunStore(tmp_path / "runs.db", "visitor-a")
    result, record = generate_and_record(
        DemoReportGenerator(), store, ProjectConfig(), DEMO_PERIOD, items
    )
    return store, record, result.report, items


def approve(store, record, report):
    store.approve(
        record.run_id,
        report,
        False,
        client_reviewed=True,
        acknowledged=[f.code for f in store.findings(record.run_id, report)],
    )


@pytest.mark.parametrize("case", CASES, ids=[c["case_id"] for c in CASES])
def test_reliability_case(case, tmp_path, monkeypatch):
    key, group, text = case["key"], case["group"], case["input"]
    items = [sample_item("A", text, source="github")]
    if group in {"normal", "messy", "context"}:
        if key == "id_collision":
            with pytest.raises(ValueError, match="conflicting"):
                normalize_evidence(items + [sample_item("A", "Different content", source="github")])
            return
        if key == "other_project":
            items[0].metadata["project_id"] = "unrelated"
            included, excluded = scoped_evidence(items, ProjectConfig())
            assert not included and excluded == items
            return
        if key == "stale":
            items[0].occurred_at = datetime(2026, 1, 1, tzinfo=UTC)
        if key == "modified_date":
            items[0].metadata["timestamp_kind"] = "document_modified"
        if key == "conflict":
            items[0].metadata["topic"] = "contract"
            items.append(sample_item("B", "Merged", topic="contract"))
        if key == "duplicate":
            assert generated(items + items) == generated(items)
        cleaned = normalize_evidence(items)
        report = generated(cleaned)
        codes = {f.code for f in validate_report(report, cleaned)}
        if key in {"quiet", "negation", "revert", "empty_description"}:
            assert report.completed == []
        elif key in {"merged", "accepted", "undeployed"}:
            assert report.completed and "does not establish deployment" in report.executive_summary
        elif key in {"ongoing", "typos", "unicode"}:
            assert report.in_progress and text.split(". ", 1)[0] in report.in_progress[0].text
        elif key == "action":
            assert report.action_items[0].owner == "PM"
            assert str(report.action_items[0].due_date) == "2026-09-04"
        elif key == "unknown_owner":
            assert {"MISSING_ACTION_OWNER", "MISSING_ACTION_DATE"} <= codes
        elif key == "multi_owner":
            assert [a.owner for a in report.action_items] == ["Dev A", "Dev B"]
        elif key == "whitespace":
            assert cleaned[0].content == "Not complete.\n\nConfirm owner."
        elif key == "conflict":
            assert "SOURCE_CONFLICT" in codes
        elif key == "stale":
            assert "OUTSIDE_PERIOD" in codes
        elif key == "dependency":
            assert "other/repo#77" in report.blockers[0].text
        elif key == "modified_date":
            assert "DOCUMENT_DATE" in codes
        elif key == "vague_title":
            assert "No description supplied" in cleaned[0].content
        return
    if group == "privacy":
        if key == "secret":
            text = "sk-ant-" + "FAKEEXAMPLE" * 4
            pem = "-----BEGIN " + "PRIVATE KEY-----\nFAKE-NOT-A-KEY\n-----END PRIVATE KEY-----"
            assert "FAKE-NOT-A-KEY" not in redact(pem)
            assert "FAKE-NOT-A-KEY" not in safe_evidence(
                [sample_item("PEM", pem)]
            )[0].model_dump_json()
        cleaned = safe_evidence([sample_item("A", text)])
        if key in {"secret", "email", "phone"}:
            assert text not in cleaned[0].model_dump_json()
            assert redact("Due 2026-09-04") == "Due 2026-09-04"
        report = generated([sample_item("A", "Merged", source="github")])
        report.executive_summary = text
        items = [sample_item("A", text)]
        if key == "private_link":
            items[0].source_url = text
        findings = validate_report(report, items)
        if key in {"secret", "email", "phone", "confidential", "blame", "private_link"}:
            assert any(f.severity == FindingSeverity.BLOCK for f in findings)
        elif key == "injection":
            assert "UNTRUSTED_INSTRUCTION" in {f.code for f in findings}
        else:
            store, record, report, _ = setup_run(tmp_path)
            with pytest.raises(ValueError, match="Review factual"):
                store.approve(
                    record.run_id,
                    report,
                    False,
                    acknowledged=[f.code for f in store.findings(record.run_id, report)],
                )
        return
    if group == "api_input":
        if key in {"malformed", "oversized"}:
            raw = text.encode() if key == "malformed" else b"x" * (2 * 1024 * 1024 + 1)
            with pytest.raises(ValueError, match="Invalid|exceeds"):
                parse_upload(raw)
        elif key == "google_tabs":
            paragraph = {"paragraph": {"elements": [{"textRun": {"content": "Not complete"}}]}}
            cell = {"content": [paragraph]}
            table = {"table": {"tableRows": [{"tableCells": [cell]}]}}
            doc = {"tabs": [{"documentTab": {"body": {"content": [table]}}}]}
            assert GoogleDocsEvidenceClient._extract_text(doc) == "Not complete"
        elif key == "google_pages":
            client = object.__new__(GoogleDocsEvidenceClient)
            client.folder_id = "allowed"
            client.drive = MagicMock()
            client.drive.files.return_value.list.return_value.execute.side_effect = [
                {"files": [{"id": "a"}], "nextPageToken": "next"},
                {"files": [{"id": "b"}]},
            ]
            assert [i["id"] for i in client.list_documents()] == ["a", "b"]
            assert client.drive.files.return_value.list.call_args.kwargs["pageToken"] == "next"
        else:
            count = 0

            def handler(request):
                nonlocal count
                count += 1
                if key == "github_pages":
                    return httpx.Response(
                        200, json=[{"number": count}] * (100 if count == 1 else 1)
                    )
                if key == "permission":
                    return httpx.Response(403)
                if count == 1:
                    if key == "timeout":
                        raise httpx.ReadTimeout("timeout", request=request)
                    return httpx.Response(429)
                return httpx.Response(200, json=[])

            with GitHubEvidenceClient("fake", "owner/repo") as client:
                client.client.close()
                client.client = httpx.Client(
                    base_url="https://example.test", transport=httpx.MockTransport(handler)
                )
                if key == "permission":
                    with pytest.raises(RuntimeError, match="denied"):
                        client._get("/repos/owner/repo/issues", {})
                    assert count == 1
                elif key == "github_pages":
                    assert len(client._paginate("/repos/owner/repo/issues", {})) == 101
                else:
                    assert client._get("/repos/owner/repo/issues", {}).status_code == 200
                    assert count == 2
        return
    store, record, report, items = setup_run(tmp_path)
    if key == "budget":
        budget = BudgetLedger(tmp_path / "budget.db", 0.1)
        budget.reserve(0.08)
        for amount in (0.03, None):
            with pytest.raises(ValueError):
                budget.reserve(amount)
        assert budget.reserved == 0.08
    elif key == "client_review":
        with pytest.raises(ValueError, match="Review factual"):
            store.approve(
                record.run_id, report, False, acknowledged=[f.code for f in record.findings]
            )
    elif key == "warning_ack":
        with pytest.raises(ValueError, match="acknowledge"):
            store.approve(record.run_id, report, False, client_reviewed=True)
    elif key == "direct_bypass":
        report.executive_summary_evidence_ids = ["INVENTED"]
        with pytest.raises(ValueError, match="blocked"):
            approve(store, record, report)
    elif key == "session":
        other = RunStore(store.path, "visitor-b")
        assert other.get(record.run_id) is None and not other.evidence(record.run_id)
        with pytest.raises(ValueError):
            other.approved_snapshot(record.run_id, report, items)
    elif key == "legacy":
        record.run_id = "legacy"
        store.save(record, report)
        assert store.get("legacy") is not None
        with pytest.raises(ValueError, match="Legacy"):
            approve(store, record, report)
    elif key in {"edit_after", "source_change"}:
        approve(store, record, report)
        if key == "edit_after":
            report.executive_summary = "Changed after approval"
        else:
            items[0].content = "Changed source"
        with pytest.raises(ValueError, match="changed"):
            approved_exports(store, record.run_id, report, items)
        assert store.get(record.run_id)[0].status != ApprovalStatus.APPROVED
    elif key == "csv_formula":
        report.action_items[0].task = text
        assert "'=SUM" in action_csv(report)
    elif key == "long_documents":
        from docx import Document

        report.executive_summary = "Café review remains in progress. " * 500
        document = Document(BytesIO(client_docx(report)))
        assert "Café" in document.paragraphs[3].text or any(
            "Café" in p.text for p in document.paragraphs
        )
        assert not any("Owner: PM" in p.text for p in document.paragraphs)
        assert client_pdf(report).startswith(b"%PDF")
    else:
        pytest.fail(f"No assertion implemented for {key}")


@pytest.mark.parametrize("variant", ["reverse", "duplicate", "spaces", "unrelated"])
def test_evidence_invariants(variant):
    items = [
        sample_item("A", "Merged audit logging.", source="github"),
        sample_item("B", "Confirm acceptance. Owner: PM."),
    ]
    original = generated(items)
    if variant == "reverse":
        items.reverse()
    elif variant == "duplicate":
        items = items + items
    elif variant == "spaces":
        items[0].content = "Merged  audit logging."
    else:
        other = sample_item("C", "Merged unrelated feature")
        other.metadata["project_id"] = "other"
        items, _ = scoped_evidence(items + [other], ProjectConfig())
    assert generated(items) == original


def test_catalog_split_is_complete():
    assert len(CASES) == 48
    assert sum(c["split"] == "held_out" for c in CASES) == 12
    assert len({c["case_id"] for c in CASES}) == 48
