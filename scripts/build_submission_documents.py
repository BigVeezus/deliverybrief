"""Build submission DOCX files from the reviewed Markdown sources.

Run with the bundled document runtime. Render with the document skill before delivery.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCES = {
    "evaluation-package.md": "DeliveryBrief-Evaluation-Package",
    "case-study.md": "DeliveryBrief-Case-Study",
    "ai-collaboration-note.md": "DeliveryBrief-AI-Collaboration-Note",
}


def table_widths(source: str, cells: list[str]) -> list[float]:
    column_count = len(cells)
    if source == "evaluation-package.md":
        header = [cell.lower() for cell in cells]
        if header == ["week", "problem pattern", "manual estimate"]:
            return [2.2, 3.45, 1.15]
        if header == ["check", "result"]:
            return [3.5, 3.3]
        if header == ["metric", "what it checks", "why it matters"]:
            return [1.8, 2.35, 2.65]
        if header == ["group", "cases", "examples"]:
            return [2.05, 0.8, 3.95]
    return [6.8 / column_count] * column_count


def plain(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text.replace("**", "").replace("`", "")


def shade_cell(cell: Any, fill: str) -> None:
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shade)


def build(source: str, name: str) -> Path:
    document = Document()
    for border in list(document.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.85)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(7)
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.line_spacing = 1.03
    table = None
    for line in (ROOT / "docs" / source).read_text().splitlines():
        if not line.strip():
            if table is not None:
                document.add_paragraph().paragraph_format.space_after = Pt(3)
            table = None
            continue
        if line.startswith("# "):
            document.add_paragraph(plain(line[2:]), "Title")
            document.add_paragraph("Prepared for final Quest submission | 8 September 2026")
        elif line.startswith("##"):
            level = len(line) - len(line.lstrip("#")) - 1
            document.add_heading(plain(line.lstrip("# ")), level=min(level, 3))
        elif line.startswith("|"):
            cells = [plain(c.strip()) for c in line.strip("|").split("|")]
            if all(re.fullmatch(r"[-: ]+", c) for c in cells):
                continue
            if table is None:
                table = document.add_table(rows=0, cols=len(cells))
                table.style = "Table Grid"
                table.autofit = False
                widths = table_widths(source, cells)
                for col, width in zip(table.columns, widths, strict=True):
                    col.width = Inches(width)
            row = table.add_row()
            for index, (cell, value) in enumerate(zip(row.cells, cells, strict=True)):
                cell.width = Inches(widths[index])
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                props = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "bottom", "left", "right"):
                    edge = OxmlElement(f"w:{side}")
                    edge.set(qn("w:val"), "single")
                    edge.set(qn("w:sz"), "4")
                    edge.set(qn("w:color"), "D9D9D9")
                    borders.append(edge)
                props.append(borders)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "bottom", "left", "right"):
                    edge = OxmlElement(f"w:{side}")
                    edge.set(qn("w:w"), "90")
                    edge.set(qn("w:type"), "dxa")
                    margins.append(edge)
                props.append(margins)
                cell.text = value
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(5)
                    header_text = table.rows[0].cells[index].text.lower() if len(table.rows) else ""
                    if header_text in {"cases", "result", "manual estimate"}:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size = Pt(9.5)
                if (
                    len(table.rows) > 1
                    and source == "evaluation-package.md"
                    and table.rows[0].cells[0].text.lower() == "check"
                    and index == 1
                ):
                    shade_cell(cell, "FFF2CC" if "$" in value else "D9EAD3")
            if len(table.rows) == 1:
                properties = row._tr.get_or_add_trPr()
                properties.append(OxmlElement("w:tblHeader"))
                for header_cell in row.cells:
                    shade_cell(header_cell, "E8EEF2")
                    for run in header_cell.paragraphs[0].runs:
                        run.bold = True
        elif line.startswith("- "):
            document.add_paragraph(plain(line[2:]), "List Bullet")
        else:
            document.add_paragraph(plain(line))
    output = ROOT / "output/docx" / f"{name}.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    return output


if __name__ == "__main__":
    for source, name in SOURCES.items():
        print(build(source, name))
