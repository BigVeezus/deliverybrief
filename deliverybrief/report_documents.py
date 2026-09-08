"""Client documents share the exact same approved report content."""

from __future__ import annotations

from io import BytesIO
from xml.sax.saxutils import escape

from deliverybrief.models import WeeklyReport


def sections(report: WeeklyReport) -> list[tuple[str, list[str]]]:
    return (
        [("Summary", [report.executive_summary])]
        + [
            (label, [i.text for i in getattr(report, field)])
            for label, field in [
                ("Completed or merged", "completed"),
                ("In progress", "in_progress"),
                ("Blockers and risks", "blockers"),
                ("Decisions", "decisions"),
                ("Next priorities", "next_priorities"),
            ]
        ]
        + [("Source limitations", report.source_limitations)]
    )


def client_docx(report: WeeklyReport) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    for name in ("Normal", "Title", "Heading 1"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(7)
    doc.add_paragraph("Weekly delivery update", "Title")
    doc.add_paragraph(report.project_name)
    doc.add_paragraph(f"Reporting period {report.period.start} to {report.period.end}")
    for heading, texts in sections(report):
        doc.add_heading(heading, level=1)
        for text in texts or ["No supported update recorded."]:
            doc.add_paragraph(text)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def client_pdf(report: WeeklyReport) -> bytes:
    from pathlib import Path

    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    output = BytesIO()
    styles = getSampleStyleSheet()
    # ReportLab ships a Unicode TrueType face, keeping local and hosted output portable.
    import reportlab

    font_path = Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"
    if "DeliveryBrief" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DeliveryBrief", str(font_path)))
    for name in ("Normal", "Title", "Heading2"):
        styles[name].fontName = "DeliveryBrief"
        styles[name].wordWrap = "CJK"
    styles["Normal"].fontSize, styles["Normal"].leading = 11, 16
    story = [
        Paragraph("Weekly delivery update", styles["Title"]),
        Paragraph(escape(report.project_name), styles["Normal"]),
        Paragraph(
            f"Reporting period {report.period.start} to {report.period.end}", styles["Normal"]
        ),
        Spacer(1, 12),
    ]
    for heading, texts in sections(report):
        story.append(Paragraph(heading, styles["Heading2"]))
        for text in texts or ["No supported update recorded."]:
            story.extend(
                [Paragraph(escape(text).replace("\n", "<br/>"), styles["Normal"]), Spacer(1, 8)]
            )
    SimpleDocTemplate(
        output, pagesize=letter, leftMargin=61, rightMargin=61, topMargin=56, bottomMargin=56
    ).build(story)
    return output.getvalue()
