from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any, cast

from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from tenacity import retry, retry_if_exception, stop_after_attempt, wait_exponential

from deliverybrief.config import parse_service_account_json
from deliverybrief.models import EvidenceItem, SourceType

SCOPES = [
    "https://www.googleapis.com/auth/drive.readonly",
    "https://www.googleapis.com/auth/documents.readonly",
]
GOOGLE_DOC_MIME = "application/vnd.google-apps.document"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT_MIMES = {
    "text/plain": "live_google_drive_text",
    "text/markdown": "live_google_drive_markdown",
    "text/csv": "live_google_drive_csv",
}
SUPPORTED_NOTE_MIMES = {GOOGLE_DOC_MIME, DOCX_MIME, *TEXT_MIMES}


class GoogleDocsIntegrationError(RuntimeError):
    """Raised when configured project notes cannot be read."""


def _transient_google_error(error: BaseException) -> bool:
    return isinstance(error, HttpError) and (error.resp.status == 429 or error.resp.status >= 500)


class GoogleDocsEvidenceClient:
    def __init__(self, service_account_json: str, folder_id: str) -> None:
        credentials = service_account.Credentials.from_service_account_info(  # type: ignore[no-untyped-call]
            parse_service_account_json(service_account_json), scopes=SCOPES
        )
        self.folder_id = folder_id
        self.drive = build("drive", "v3", credentials=credentials, cache_discovery=False)
        self.docs = build("docs", "v1", credentials=credentials, cache_discovery=False)

    @retry(
        retry=retry_if_exception(_transient_google_error),
        stop=stop_after_attempt(3),
        wait=wait_exponential(multiplier=0.5, min=0.5, max=4),
        reraise=True,
    )
    def list_documents(self) -> list[dict[str, str]]:
        try:
            records: list[dict[str, str]] = []
            page_token = None
            mime_query = " or ".join(
                f"mimeType = '{mime_type}'" for mime_type in sorted(SUPPORTED_NOTE_MIMES)
            )
            for _ in range(20):
                response = (
                    self.drive.files()
                    .list(
                        q=(
                            f"'{self.folder_id}' in parents and trashed = false and "
                            f"({mime_query})"
                        ),
                        fields="nextPageToken,files(id,name,mimeType,modifiedTime,webViewLink)",
                        orderBy="modifiedTime desc",
                        pageSize=100,
                        pageToken=page_token,
                    )
                    .execute()
                )
                records.extend(response.get("files", []))
                page_token = response.get("nextPageToken")
                if not page_token:
                    return records
            raise GoogleDocsIntegrationError(
                "Folder exceeds the document limit. Use a smaller folder."
            )
        except HttpError as error:
            if error.resp.status in {401, 403, 404}:
                raise GoogleDocsIntegrationError(
                    "Google denied access. Share the configured folder with the service account "
                    "as Viewer."
                ) from error
            raise

    @retry(
        retry=retry_if_exception(_transient_google_error),
        stop=stop_after_attempt(3),
        wait=wait_exponential(multiplier=0.5, min=0.5, max=4),
        reraise=True,
    )
    def _document(self, document_id: str) -> dict[str, Any]:
        try:
            response = (
                self.docs.documents().get(documentId=document_id, includeTabsContent=True).execute()
            )
            return cast(dict[str, Any], response)
        except HttpError as error:
            if error.resp.status in {401, 403, 404}:
                raise GoogleDocsIntegrationError(
                    f"Google document {document_id} is unavailable to the service account."
                ) from error
            raise

    @retry(
        retry=retry_if_exception(_transient_google_error),
        stop=stop_after_attempt(3),
        wait=wait_exponential(multiplier=0.5, min=0.5, max=4),
        reraise=True,
    )
    def _download_file(self, file_id: str) -> bytes:
        try:
            content = self.drive.files().get_media(fileId=file_id).execute()
        except HttpError as error:
            if error.resp.status in {401, 403, 404}:
                raise GoogleDocsIntegrationError(
                    f"Google Drive file {file_id} is unavailable to the service account."
                ) from error
            raise
        if isinstance(content, bytes):
            return content
        return str(content or "").encode()

    def _text_file(self, file_id: str) -> str:
        return self._download_file(file_id).decode("utf-8", errors="replace").strip()

    @staticmethod
    def _extract_text(document: dict[str, Any]) -> str:
        parts: list[str] = []

        def walk(node: Any) -> None:
            if isinstance(node, list):
                for child in node:
                    walk(child)
            elif isinstance(node, dict):
                if "textRun" in node:
                    parts.append(str(node["textRun"].get("content", "")))
                else:
                    for key in (
                        "body",
                        "content",
                        "paragraph",
                        "elements",
                        "table",
                        "tableRows",
                        "tableCells",
                        "documentTab",
                        "childTabs",
                    ):
                        if key in node:
                            walk(node[key])
                    if "paragraph" in node:
                        parts.append("\n")

        walk(document.get("tabs") or document)
        return "".join(parts).strip()

    @staticmethod
    def _extract_docx_text(raw: bytes) -> str:
        from docx import Document

        document = Document(BytesIO(raw))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if values:
                    paragraphs.append(" | ".join(values))
        return "\n".join(paragraphs).strip()

    def collect(self, selected_document_ids: list[str]) -> list[EvidenceItem]:
        available = {item["id"]: item for item in self.list_documents()}
        evidence: list[EvidenceItem] = []
        for document_id in selected_document_ids:
            metadata = available.get(document_id)
            if not metadata:
                raise GoogleDocsIntegrationError(
                    f"Selected document {document_id} is not in the configured folder."
                )
            mime_type = metadata.get("mimeType")
            if mime_type == GOOGLE_DOC_MIME:
                document = self._document(document_id)
                content = self._extract_text(document)
                evidence_id = f"GDOC-{document_id}"
                ingestion = "live_google_docs"
            elif mime_type in TEXT_MIMES:
                content = self._text_file(document_id)
                evidence_id = f"GDRIVE-TEXT-{document_id}"
                ingestion = TEXT_MIMES[mime_type]
            elif mime_type == DOCX_MIME:
                content = self._extract_docx_text(self._download_file(document_id))
                evidence_id = f"GDRIVE-DOCX-{document_id}"
                ingestion = "live_google_drive_docx"
            else:
                raise GoogleDocsIntegrationError(
                    f"Unsupported note type for {metadata['name']}: {mime_type}"
                )
            if not content:
                raise GoogleDocsIntegrationError(
                    f"Document {metadata['name']} contains no readable text."
                )
            modified = datetime.fromisoformat(metadata["modifiedTime"].replace("Z", "+00:00"))
            evidence.append(
                EvidenceItem(
                    evidence_id=evidence_id,
                    source=SourceType.GOOGLE_DOC,
                    title=metadata["name"],
                    content=content,
                    occurred_at=modified,
                    source_url=metadata.get("webViewLink"),
                    metadata={
                        "kind": "project_note",
                        "document_id": document_id,
                        "mime_type": mime_type,
                        "timestamp_kind": "document_modified",
                        "ingestion": ingestion,
                    },
                )
            )
        return evidence
